import requests
import pandas as pd
import json
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import numbers

def scrape_and_create_report():
    url = "https://ackodrive.com/collection/maruti-suzuki-cars/"
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"
    }
    response = requests.get(url, headers=headers)
    soup = BeautifulSoup(response.text, 'html.parser')
    script = soup.find('script', id='__NEXT_DATA__')
    data = json.loads(script.string)
    cars = data['props']['pageProps']['listingData']['result']
    variants = []
    for model in cars:
        for variant in model.get('variants', []):
            price = variant.get('price', {}).get('final_price', 0)
            basic = variant.get('basic_feature', {})
            colors = [c['color']['brand_color'] for c in variant.get('colors', [])]
            variants.append({
                "Make": model.get('brand_name', ''),
                "Model": model.get('model_name', ''),
                "Variant": variant.get('variant_name', ''),
                "Price": price,
                "Fuel Type": basic.get('fuel_type', ''),
                "Transmission": basic.get('transmission', ''),
                "Engine CC": basic.get('engine', ''),
                "Available Colors": ', '.join(sorted(colors))
            })
    df = pd.DataFrame(variants)
    file = 'maruti_suzuki_cari3_report.xlsx'
    with pd.ExcelWriter(file, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='All Variants')
        summary = df.groupby('Model')['Price'].mean().round(0).reset_index()
        summary.columns = ['Model', 'Average Price']
        summary.to_excel(writer, index=False, sheet_name='Dashboard')
    wb = load_workbook(file)
    vsheet = wb['All Variants']
    for cell in vsheet['D']:
        cell.number_format = '₹#,##0'
    dsheet = wb['Dashboard']
    for cell in dsheet['B']:
        cell.number_format = '₹#,##0'
    chart = BarChart()
    chart.title = "Average Price by Car Model"
    chart.y_axis.title = "Average Price (₹)"
    chart.x_axis.title = "Car Model"
    chart.legend = None
    rows = len(summary)
    data = Reference(dsheet, min_col=2, min_row=1, max_row=rows + 1)
    cats = Reference(dsheet, min_col=1, min_row=2, max_row=rows + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 15
    chart.width = 30
    dsheet.add_chart(chart, "D2")
    wb.save(file)

if __name__ == '__main__':
    scrape_and_create_report()
